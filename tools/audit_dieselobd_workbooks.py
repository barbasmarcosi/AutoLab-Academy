"""Read-only, provenance-pinned structural inspection of the DieselOBD source workbooks.

Requires openpyxl (installed only by the research workflow). No training, vehicle
health labels, invented timestamps or redistribution of upstream spreadsheet data.
"""
import argparse
import hashlib
import json
import math
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

EXPECTED = {
    "MASTER_TRAIN_X.xlsx": ("ca5fcf99ae8d2edc7c04b5636a1b18839df7dd52", 8995072),
    "MASTER_TEST_X.xlsx": ("f76b794fac48ee3f2ad74214757272e7ab2cd9e6", 2531527),
}
PIDS = ("LOAD_PCT","ECT","MAP","RPM","VSS","IAT","MAF","FRP","BARO","VPWR","AAT")
VEHICLE_KEYS = {"vehicle_id","vehicle","vehicleid","car_id"}
DTC_KEYS = {"P0000","P0562","P0113","P0102","P0403","P0404","P2562","P2015","P2009","P0107","P0069","P0089","P0234","P0406"}

def blob_sha1(path):
    size = path.stat().st_size
    digest = hashlib.sha1()
    digest.update(f"blob {size}\0".encode("ascii"))
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()

def audit_workbook(path):
    path = Path(path)
    if path.name not in EXPECTED:
        raise ValueError("Unexpected upstream workbook name")
    expected_blob, expected_bytes = EXPECTED[path.name]
    if path.stat().st_size != expected_bytes or blob_sha1(path) != expected_blob:
        raise ValueError("DieselOBD upstream source integrity mismatch")
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    results = []
    try:
        for sheet in wb.worksheets:
            rows = sheet.iter_rows(values_only=True)
            head = next(rows, None)
            if head is None:
                results.append({"sheet":sheet.title,"rows":0,"error":"empty sheet"})
                continue
            headers = [str(v).strip() if v is not None else "" for v in head]
            normalized = [x.upper() for x in headers]
            vehicle_index = next((i for i,h in enumerate(headers) if h.lower() in VEHICLE_KEYS),None)
            time_columns = [h for h in headers if h.lower() in {"timestamp","time_ms","time","datetime"}]
            pids = [p for p in PIDS if p in normalized]
            dtcs = [h for h in headers if h.upper() in DTC_KEYS]
            counts = Counter()
            present = Counter()
            invalid = Counter()
            unit_set = set()
            row_count = 0
            for row in rows:
                if not row or all(v is None for v in row):
                    continue
                row_count += 1
                if vehicle_index is not None and vehicle_index < len(row) and row[vehicle_index] is not None:
                    unit_set.add(str(row[vehicle_index]).strip())
                for j,h in enumerate(headers):
                    if not h:
                        continue
                    v = row[j] if j < len(row) else None
                    if v is None or v == "":
                        continue
                    present[h] += 1
                    if h in dtcs and str(v).strip() not in {"0","0.0","False","false"}:
                        counts[h] += 1
                    if h.upper() in PIDS:
                        try:
                            if not math.isfinite(float(v)):
                                invalid[h] += 1
                        except (ValueError,TypeError):
                            invalid[h] += 1
            results.append({
                "sheet":sheet.title,"rows":row_count,"columnCount":len(headers),
                "headers":headers,"obdPidColumns":pids,"dtcColumns":dtcs,
                "timestampColumns":time_columns,
                "vehicleIdColumn":headers[vehicle_index] if vehicle_index is not None else None,
                "distinctVehicleCount":len(unit_set) if vehicle_index is not None else None,
                "missingFraction":{h:round(1-present[h]/row_count,6) for h in headers if h and row_count},
                "invalidPidCount":dict(invalid),
                "dtcPositiveRows":dict(counts),
                "usableForAging":False,
                "limits":"Fault labels are not direct wear measures; do not claim vehicle independence if no vehicle ID."
            })
    finally:
        wb.close()
    return {"schema":"MA00-dieselobd-xlsx-audit-v1","sourceFile":path.name,
            "sourceGitBlob":expected_blob,"sourceBytes":expected_bytes,
            "sheets":results,"wearGroundTruthPresent":False,
            "eligibleForMechanicalAgeValidation":False}

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("train")
    parser.add_argument("test")
    args=parser.parse_args()
    result=[audit_workbook(args.train),audit_workbook(args.test)]
    print(json.dumps(result,ensure_ascii=False,indent=2,allow_nan=False))

if __name__=="__main__":
    main()
